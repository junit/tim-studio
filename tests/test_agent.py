# -*- coding: utf-8 -*-
"""
AI 智能体系统单元测试：BOT 实体管理、提示词自我增强、长期记忆持久化、技能自沉淀与 API 路由
"""

import os
import sys
import shutil
import time
from unittest.mock import patch
import pytest

# 确保加载项目根路径
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from fastapi.testclient import TestClient

from app import app
from core.agent.bot_manager import BotManager
from core.agent.tools import RunSkillTool, build_bot_tools
from core.tim_client import TimClient


@pytest.fixture
def temp_bot_manager(tmp_path):
    """创建隔离的临时 BotManager"""
    data_dir = str(tmp_path / "data")
    mgr = BotManager(data_dir=data_dir)
    return mgr


def test_default_bot_created(temp_bot_manager):
    """验证初始化时自动创建默认助手"""
    bots = temp_bot_manager.list_bots()
    assert len(bots) >= 1
    bot = temp_bot_manager.get_bot("default_assistant")
    assert bot is not None
    assert bot.name == "TIM 全能助手"
    assert "166" in bot.description


def test_custom_bot_lifecycle(temp_bot_manager):
    """测试自定义 BOT 创建、更新与删除"""
    bot = temp_bot_manager.create_bot(
        bot_id="security_officer",
        name="社群风控官",
        avatar="🛡️",
        description="专注敏感发言拦截与禁言",
        base_prompt="你是一个严谨的 IM 安全合规助手。"
    )
    assert bot.id == "security_officer"
    assert bot.avatar == "🛡️"

    # 查询
    retrieved = temp_bot_manager.get_bot("security_officer")
    assert retrieved is not None
    assert retrieved.name == "社群风控官"

    # 更新
    retrieved.name = "社群超级风控官"
    temp_bot_manager.save_bot(retrieved)
    updated = temp_bot_manager.get_bot("security_officer")
    assert updated.name == "社群超级风控官"

    # 删除
    ok = temp_bot_manager.delete_bot("security_officer")
    assert ok is True
    assert temp_bot_manager.get_bot("security_officer") is None


def test_prompt_self_enhancement(temp_bot_manager):
    """测试 BOT 提示词自我增强 (Prompt Self-Enhancement)"""
    bot_id = "default_assistant"
    guideline = temp_bot_manager.add_learned_guideline(
        bot_id=bot_id,
        instruction="对于所有禁言操作，若未显式指定禁言时间，默认按 3600 秒（1小时）处理。",
        reason="用户在历史操作中要求统一默认禁言时长"
    )
    assert guideline is not None
    assert guideline.id.startswith("g-")

    bot = temp_bot_manager.get_bot(bot_id)
    assert len(bot.learned_guidelines) == 1
    assert "3600 秒" in bot.learned_guidelines[0].instruction
    assert "统一默认" in bot.learned_guidelines[0].reason

    # 删除准则
    removed = temp_bot_manager.remove_learned_guideline(bot_id, guideline.id)
    assert removed is True
    bot_after = temp_bot_manager.get_bot(bot_id)
    assert len(bot_after.learned_guidelines) == 0


def test_long_term_memory_persistence(temp_bot_manager):
    """测试长期事实记忆系统的存取与删除"""
    bot_id = "default_assistant"

    # 保存记忆
    item = temp_bot_manager.save_memory(bot_id, "primary_test_group", "@TGS#12345678")
    assert item.key == "primary_test_group"
    assert item.content == "@TGS#12345678"

    # 查阅列表
    memories = temp_bot_manager.list_memories(bot_id)
    assert len(memories) == 1
    assert memories[0].key == "primary_test_group"

    # 更新已存在的记忆
    item2 = temp_bot_manager.save_memory(bot_id, "primary_test_group", "@TGS#87654321")
    assert item2.content == "@TGS#87654321"
    assert len(temp_bot_manager.list_memories(bot_id)) == 1

    # 删除记忆
    del_ok = temp_bot_manager.delete_memory(bot_id, "primary_test_group")
    assert del_ok is True
    assert len(temp_bot_manager.list_memories(bot_id)) == 0


def test_skill_self_creation_and_execution(temp_bot_manager):
    """测试技能自编写、保存与沙箱执行"""
    bot_id = "default_assistant"
    skill_code = """
def execute(client, **kwargs):
    user = kwargs.get("user", "test_user")
    multiplier = kwargs.get("multiplier", 2)
    return {"processed_user": user.upper(), "score": multiplier * 10}
"""
    # 保存技能
    skill = temp_bot_manager.save_skill(
        bot_id=bot_id,
        name="calculate_user_score",
        description="测试计算用户积分",
        code=skill_code,
        input_schema={"user": "用户名", "multiplier": "倍数"}
    )
    assert skill.name == "calculate_user_score"

    skills = temp_bot_manager.list_skills(bot_id)
    assert len(skills) == 1
    assert skills[0].name == "calculate_user_score"

    # 执行技能
    client = TimClient()
    runner = RunSkillTool(bot_id, temp_bot_manager, client)
    res = runner.forward("calculate_user_score", {"user": "alice", "multiplier": 3})
    assert res.get("success") is True
    assert res.get("result", {}).get("processed_user") == "ALICE"
    assert res.get("result", {}).get("score") == 30

    # 删除技能
    del_ok = temp_bot_manager.delete_skill(bot_id, "calculate_user_score")
    assert del_ok is True
    assert len(temp_bot_manager.list_skills(bot_id)) == 0


def test_agent_api_routes():
    """测试 Agent 后端 REST API 接口"""
    client = TestClient(app)

    # 1. 配置获取
    cfg_resp = client.get("/api/agent/config")
    assert cfg_resp.status_code == 200
    cfg_data = cfg_resp.json()
    assert "providers" in cfg_data
    assert "config" in cfg_data

    # 2. BOT 列表
    bots_resp = client.get("/api/agent/bots")
    assert bots_resp.status_code == 200
    assert len(bots_resp.json().get("bots", [])) >= 1

    # 2.1 验证内置智能体受保护不可修改基础人设
    default_up_resp = client.put("/api/agent/bots/default_assistant", json={
        "base_prompt": "尝试篡改内置人设"
    })
    assert default_up_resp.status_code == 400
    assert "内置智能体基础设定受系统保护" in default_up_resp.json()["detail"]

    # 3. 创建临时测试 BOT
    test_id = f"test_bot_{os.getpid()}"
    create_resp = client.post("/api/agent/bots", json={
        "id": test_id,
        "name": "接口自动化 BOT",
        "avatar": "⚡",
        "description": "API 自动化测试专属",
        "base_prompt": "你是一个测试专属助理"
    })
    assert create_resp.status_code == 200

    # 3.1 获取 BOT 详情并验证组装的系统提示词
    detail_resp = client.get(f"/api/agent/bots/{test_id}")
    assert detail_resp.status_code == 200
    detail_data = detail_resp.json().get("bot", {})
    assert "assembled_prompt" in detail_data
    assert "你是一个测试专属助理" in detail_data["assembled_prompt"]
    assert "自主进化与无感沉淀机制" in detail_data["assembled_prompt"]

    # 3.2 更新 BOT 基础人设并验证动态重装
    update_resp = client.put(f"/api/agent/bots/{test_id}", json={
        "base_prompt": "你是一个已更新的高级自动化测试专家"
    })
    assert update_resp.status_code == 200
    updated_bot = update_resp.json().get("bot", {})
    assert "assembled_prompt" in updated_bot
    assert "你是一个已更新的高级自动化测试专家" in updated_bot["assembled_prompt"]

    # 4. 记忆操作
    mem_resp = client.post(f"/api/agent/bots/{test_id}/memory", json={
        "key": "app_env",
        "content": "staging"
    })
    assert mem_resp.status_code == 200
    get_mem = client.get(f"/api/agent/bots/{test_id}/memory")
    assert len(get_mem.json().get("memories", [])) == 1

    # 5. 自增强准则操作
    guide_resp = client.post(f"/api/agent/bots/{test_id}/guidelines", json={
        "instruction": "测试环境下每次修改前打印日志",
        "reason": "审计要求"
    })
    assert guide_resp.status_code == 200
    guide_id = guide_resp.json()["guideline"]["id"]

    del_guide = client.delete(f"/api/agent/bots/{test_id}/guidelines/{guide_id}")
    assert del_guide.status_code == 200

    # 6. 用户/前端一键沉淀技能操作
    skill_resp = client.post(f"/api/agent/bots/{test_id}/skills", json={
        "name": "auto_test_skill",
        "description": "由前端/用户一键沉淀的技能",
        "python_code": "def execute(client, **kwargs): return {'ok': True}"
    })
    assert skill_resp.status_code == 200
    skills_list = client.get(f"/api/agent/bots/{test_id}/skills")
    assert len(skills_list.json().get("skills", [])) == 1

    # 7. 删除临时 BOT
    del_resp = client.delete(f"/api/agent/bots/{test_id}")
    assert del_resp.status_code == 200


def test_auto_generate_bot_profile():
    """测试通过仅输入 BOT 名字和目标，AI 自动生成 ID、头像与提示词"""
    client = TestClient(app)

    # 1. 测试生成预览接口
    preview_resp = client.post("/api/agent/bots/generate-preview", json={
        "name": "社群风控卫士",
        "goal": "监控群内敏感广告发言，识别刷屏成员并自动禁言"
    })
    assert preview_resp.status_code == 200
    preview_data = preview_resp.json()
    assert preview_data["success"] is True
    profile = preview_data["profile"]
    assert profile["avatar"] == "🛡️"
    assert "社群风控卫士" in profile["base_prompt"]
    assert "166" in profile["base_prompt"]

    # 2. 测试仅填名字和目标，一键自动创建 BOT
    create_resp = client.post("/api/agent/bots", json={
        "name": "消息巡检大师",
        "goal": "每日巡检重点大群消息流，统计未读与违规撤回情况"
    })
    assert create_resp.status_code == 200
    bot_data = create_resp.json()["bot"]
    assert bot_data["name"] == "消息巡检大师"
    assert bot_data["id"].startswith("bot_")
    assert len(bot_data["avatar"]) > 0
    assert len(bot_data["base_prompt"]) > 50

    # 清理
    del_resp = client.delete(f"/api/agent/bots/{bot_data['id']}")
    assert del_resp.status_code == 200


def test_history_runtime_metadata_persistence(temp_bot_manager):
    """测试会话历史中思维链、Python代码、执行日志与工具调用的完整持久化与读取"""
    bot_id = "default_assistant"

    # 1. 存入包含丰富运行期元数据的 assistant 消息
    metadata_payload = {
        "thought": "正在思考如何处理群禁言...",
        "code_blocks": [{"step": 1, "code": "res = tim_call('set_group_mute', group_id='@TGS#1')"}],
        "execution_logs": [{"step": 1, "output": "Successfully executed sandbox task"}],
        "tool_calls": [{"step": 1, "output": "{'ErrorCode': 0, 'ActionStatus': 'OK'}"}],
        "code_count": 1,
        "model": "deepseek-chat"
    }

    temp_bot_manager.append_history(
        bot_id=bot_id,
        role="assistant",
        content="已成功设置群禁言。",
        metadata=metadata_payload
    )

    # 2. 读取历史，断言元数据完整性
    history = temp_bot_manager.get_history(bot_id)
    assert len(history) >= 1
    last_msg = history[-1]
    assert last_msg["role"] == "assistant"
    assert last_msg["content"] == "已成功设置群禁言。"
    assert "metadata" in last_msg
    meta = last_msg["metadata"]
    assert meta["thought"] == "正在思考如何处理群禁言..."
    assert len(meta["code_blocks"]) == 1
    assert "set_group_mute" in meta["code_blocks"][0]["code"]
    assert len(meta["execution_logs"]) == 1
    assert "Successfully executed" in meta["execution_logs"][0]["output"]
    assert len(meta["tool_calls"]) == 1


def test_history_pagination_and_atomic_writes(temp_bot_manager):
    """测试会话历史的分页查询（before_id / has_more）与原子写入安全性"""
    bot_id = "default_assistant"
    # 模拟写入 10 条消息 (5 轮对话)
    for i in range(10):
        temp_bot_manager.append_history(bot_id, "user" if i % 2 == 0 else "assistant", f"msg_{i}")

    # 1. 查询默认最新批次 (limit=4)
    paged = temp_bot_manager.get_history_paged(bot_id, limit=4)
    assert paged["total"] >= 10
    assert len(paged["history"]) == 4
    assert paged["has_more"] is True
    assert paged["history"][-1]["content"] == "msg_9"
    first_id = paged["history"][0]["id"]

    # 2. 向上分页查询更早历史 (before_id)
    paged_prev = temp_bot_manager.get_history_paged(bot_id, limit=4, before_id=first_id)
    assert len(paged_prev["history"]) == 4
    assert paged_prev["history"][-1]["content"] == "msg_5"
    assert paged_prev["has_more"] is True


def test_format_execution_observation_markdown():
    """测试沙箱执行输出向结构化 Markdown 的转换"""
    from core.agent.engine import format_execution_observation

    # 1. 字典转换为 ```json 格式
    res_dict = {"ActionStatus": "OK", "ErrorCode": 0, "MsgSeq": 100}
    md_dict = format_execution_observation(res_dict)
    assert md_dict.startswith("```json")
    assert '"ActionStatus": "OK"' in md_dict

    # 2. 带有 Print outputs: 前缀的 Python 字典字符串
    raw_str = "Print outputs:\n{'ActionStatus': 'OK'}"
    md_str = format_execution_observation(raw_str)
    assert md_str.startswith("```json")

    # 3. 错误信息转换为告警引用块
    err_str = "Traceback (most recent call last):\nValueError: invalid argument"
    md_err = format_execution_observation(err_str)
    assert "> ⚠️ **沙箱执行告警 / 异常**" in md_err


def test_agent_streaming_event_flow(temp_bot_manager):
    """测试真正的流式输出事件序列（start -> thought_chunk -> code_start -> execution_result -> answer_start -> chunk -> done）"""
    from unittest.mock import MagicMock, patch
    from core.agent.engine import agent_engine
    from smolagents.models import ChatMessageStreamDelta
    from smolagents.memory import ToolCall, ActionStep, FinalAnswerStep, Timing
    from smolagents.agents import ActionOutput

    mock_agent = MagicMock()

    # 模拟真实流式执行中的事件流
    mock_steps = [
        ChatMessageStreamDelta(content="<think>正在规划群组查询路径...</think>"),
        ToolCall(name="python_interpreter", arguments="res = tim_call('get_joined_group_list')", id="call_1"),
        ActionOutput(output={"ActionStatus": "OK", "ErrorCode": 0, "TotalCount": 5}, is_final_answer=False),
        ActionStep(step_number=1, timing=Timing(start_time=0.0)),
        FinalAnswerStep(output="已为您成功查询到 5 个群组。")
    ]
    mock_agent.run.return_value = iter(mock_steps)

    with patch("core.agent.engine.CodeAgent", return_value=mock_agent):
        with patch.object(agent_engine, "create_model", return_value=MagicMock()):
            with patch.object(agent_engine, "manager", temp_bot_manager):
                events = list(agent_engine.run_agent_stream("default_assistant", "查群"))
                event_names = [e["event"] for e in events]

            # 验证全流式事件完整性
            assert "start" in event_names
            assert "thought_chunk" in event_names
            assert "code_start" in event_names
            assert "execution_result" in event_names
            assert "answer_start" in event_names
            assert "chunk" in event_names
            assert "message" in event_names
            assert "done" in event_names


def test_clean_prompt_text():
    """测试清理 ChatMessage 调试对象 repr 字符串"""
    raw = (
        "ChatMessage(role='assistant', "
        "content='你是【消息管理员】腾讯云 IM 智能体，以专业、严谨、可审计的技术人设，协助管理与排查消息管理相关 IM 业务需求。', "
        "tool_calls=None, raw=ChatCompletion(id='test-123'))"
    )
    cleaned = BotManager.clean_prompt_text(raw)
    assert not cleaned.startswith("ChatMessage(")
    assert cleaned.startswith("你是【消息管理员】")
    assert "tool_calls" not in cleaned


def test_infer_suggested_prompts():
    """测试根据不同人格推断个性化快捷问题"""
    # 风控
    prompts_sec = BotManager.infer_suggested_prompts("社群风控官", "排查违规刷屏并禁言")
    assert any("禁言" in p or "黑名单" in p for p in prompts_sec)

    # 巡检
    prompts_ins = BotManager.infer_suggested_prompts("IM巡检专家", "监控接口健康度与异常离线")
    assert any("巡检" in p or "监控" in p for p in prompts_ins)

    # 投研
    prompts_fin = BotManager.infer_suggested_prompts("投资研究员", "管理投研群与研报分发")
    assert any("投研" in p or "研报" in p for p in prompts_fin)


def test_delete_bot_api_restrictions(temp_bot_manager):
    """测试 API 路由删除智能体及对默认助手的安全保护"""
    client = TestClient(app)

    # 1. 创建自定义 bot
    temp_bot_manager.create_bot(
        bot_id="to_be_deleted",
        name="临时助手",
        avatar="🧪",
        description="待删除测试",
        base_prompt="临时设定"
    )

    # 2. 尝试删除默认助手应被拒绝 (400)
    res_default = client.delete("/api/agent/bots/default_assistant")
    assert res_default.status_code == 400
    assert "无法删除" in res_default.json()["detail"]


def test_agent_run_manager_and_active_run_api():
    """测试后台会话管理器及 active-run 查询与停止接口"""
    from core.agent.run_manager import AgentRun, AgentRunManager

    mgr = AgentRunManager()
    run = AgentRun("test_run_1", "default_assistant", "测试后台运行", client_msg_id="ast-custom-123")
    assert run.assistant_msg_id == "ast-custom-123"
    assert run.status == "running"
    assert run.stream_status == "thinking"

    # 模拟事件分发
    run.dispatch_event("thought_chunk", {"delta": "思考中..."})
    assert run.thought == "思考中..."
    run.dispatch_event("code_start", {"step": 1, "code": "print('ok')"})
    assert run.stream_status == "executing"
    assert len(run.code_blocks) == 1
    assert run.code_blocks[0]["status"] == "running"

    run.dispatch_event("execution_result", {"step": 1, "output": "ok"})
    assert run.code_blocks[0]["status"] == "success"
    assert len(run.execution_logs) == 1

    # 测试快照
    snap = run.to_dict()
    assert snap["run_id"] == "test_run_1"
    assert snap["thought"] == "思考中..."
    assert len(snap["code_blocks"]) == 1

    # 测试 API: 无进行中任务
    client = TestClient(app)
    res = client.get("/api/agent/bots/default_assistant/active-run")
    assert res.status_code == 200
    assert "has_active_run" in res.json()

    # 测试停止接口
    stop_res = client.post("/api/agent/bots/default_assistant/stop")
    assert stop_res.status_code == 200
    assert "success" in stop_res.json()


def test_builtin_bot_exclusive_create_bot_tool(temp_bot_manager):
    """测试仅内建全能助手拥有 create_bot 权限，且创建后可持久化"""
    from core.agent.tools import build_bot_tools
    from core.tim_client import TimClient

    tim = TimClient(config={"sdkappid": 1400000000, "admin_user": "adm", "secret_key": "key"})

    # 1. 内建助手拥有 create_bot
    tools_default = build_bot_tools("default_assistant", temp_bot_manager, tim)
    tool_names_default = [t.name for t in tools_default]
    assert "create_bot" in tool_names_default
    assert "list_agents" in tool_names_default
    assert "call_agent" in tool_names_default

    # 2. 自定义 BOT 严禁拥有 create_bot
    tools_custom = build_bot_tools("my_custom_bot", temp_bot_manager, tim)
    tool_names_custom = [t.name for t in tools_custom]
    assert "create_bot" not in tool_names_custom
    assert "list_agents" in tool_names_custom
    assert "call_agent" in tool_names_custom

    # 3. 执行 create_bot 工具创建新智能体
    create_tool = next(t for t in tools_default if t.name == "create_bot")
    res = create_tool.forward(
        name="自动化巡检官",
        goal="监控群消息并发与异常掉线",
        avatar="🔍"
    )
    assert res["success"] is True
    assert "自动化巡检官" in res["message"]
    created_id = res["bot"]["id"]
    persisted = temp_bot_manager.get_bot(created_id)
    assert persisted is not None
    assert persisted.name == "自动化巡检官"
    assert persisted.avatar == "🔍"


def test_multi_agent_coordination_and_clean_delete(temp_bot_manager):
    """测试智能体协同调度 (CallAgent) 及安全彻底删除数据"""
    from core.agent.tools import CallAgentTool, ListAgentsTool

    # 1. 测试 ListAgentsTool
    list_tool = ListAgentsTool(temp_bot_manager)
    agents = list_tool.forward()
    assert len(agents) >= 1
    assert any(a["id"] == "default_assistant" for a in agents)

    # 2. 创建协同目标智能体
    target_bot = temp_bot_manager.create_bot(
        bot_id="security_bot",
        name="社群风控官",
        avatar="🛡️",
        description="风控与违规禁言",
        base_prompt="专业风控人设"
    )

    # 3. 测试 CallAgentTool 自调用保护
    call_tool = CallAgentTool("security_bot", temp_bot_manager)
    self_res = call_tool.forward("@社群风控官", "自我调用")
    assert "自身" in self_res

    # 4. 模拟 mock 协同调度：支持无 @、括号、自然口语简称等各种自然语言匹配
    from unittest.mock import MagicMock
    mock_engine = MagicMock()
    mock_engine.execute_subtask.return_value = "已完成敏感词检查，未发现违规发言。"

    caller_tool = CallAgentTool("default_assistant", temp_bot_manager, engine=mock_engine)

    # 4.1 纯文本名称（无 @）
    coord_res1 = caller_tool.forward("社群风控官", "排查违规")
    assert "社群风控官" in coord_res1

    # 4.2 口语简称与核心词（如“风控”）
    coord_res2 = caller_tool.forward("风控", "排查违规")
    assert "社群风控官" in coord_res2

    # 4.3 携带括号或特殊符号
    coord_res3 = caller_tool.forward("【社群风控官】(security_bot)", "排查违规")
    assert "社群风控官" in coord_res3

    assert mock_engine.execute_subtask.call_count == 3

    # 5. 测试彻底清理删除
    b_dir = temp_bot_manager._bot_dir("security_bot")
    assert os.path.exists(b_dir)
    assert temp_bot_manager.delete_bot("security_bot") is True
    assert not os.path.exists(b_dir)
    assert temp_bot_manager.get_bot("security_bot") is None


def test_active_runs_api():
    """测试多智能体活跃运行状态查询接口及 run_manager 跟踪"""
    from core.agent.run_manager import agent_run_manager, AgentRun
    client = TestClient(app)

    # 1. 初始查询活跃状态
    resp = client.get("/api/agent/active-runs")
    assert resp.status_code == 200
    data = resp.json()
    assert "active_runs" in data

    # 2. 注入模拟的活跃任务
    test_run = AgentRun("run_test_bot_1", "test_bot_1", "请检查当前所有群组", client_msg_id="ast-1")
    test_run.stream_status = "executing"
    test_run.current_step = 2

    with agent_run_manager._lock:
        agent_run_manager.runs["run_test_bot_1"] = test_run
        agent_run_manager.active_bot_runs["test_bot_1"] = "run_test_bot_1"

    # 3. 再次查询接口，确认包含 test_bot_1
    resp2 = client.get("/api/agent/active-runs")
    assert resp2.status_code == 200
    active_runs = resp2.json()["active_runs"]
    assert "test_bot_1" in active_runs
    bot_run = active_runs["test_bot_1"]
    assert bot_run["status"] == "running"
    assert bot_run["stream_status"] == "executing"
    assert bot_run["current_step"] == 2
    assert "请检查当前所有群组" in bot_run["user_message"]

    # 4. 任务结束或停止后，状态清理
    test_run.status = "done"
    with agent_run_manager._lock:
        agent_run_manager.active_bot_runs.pop("test_bot_1", None)

    resp3 = client.get("/api/agent/active-runs")
    assert resp3.status_code == 200
    assert "test_bot_1" not in resp3.json()["active_runs"]


def test_upload_file_and_read_tool(temp_bot_manager):
    """测试多格式文件上传、智能解析摘要以及 ReadFileTool 读取能力"""
    import io
    from core.agent.tools import ReadFileTool
    client = TestClient(app)

    # 1. 测试 CSV 文件上传与结构解析
    csv_content = b"user_id,nickname,role\nadmin_01,\xe7\xae\xa1\xe7\x90\x86\xe5\x91\x98,admin\nuser_02,\xe6\xb5\x8b\xe8\xaf\x95\xe5\x91\x98,member\n"
    csv_file = io.BytesIO(csv_content)
    csv_resp = client.post(
        "/api/agent/bots/default_assistant/upload",
        files={"file": ("users.csv", csv_file, "text/csv")}
    )
    assert csv_resp.status_code == 200
    csv_data = csv_resp.json()
    assert csv_data["success"] is True
    file_info = csv_data["file"]
    assert file_info["type"] == "csv"
    assert file_info["category"] == "table"
    assert "user_id" in file_info["summary"]
    assert os.path.exists(file_info["path"])

    # 2. 测试文件在线读取接口
    dl_resp = client.get(file_info["url"])
    assert dl_resp.status_code == 200
    assert b"user_id" in dl_resp.content

    # 3. 测试 ReadFileTool 工具
    read_tool = ReadFileTool("default_assistant", temp_bot_manager)
    tool_out = read_tool.forward(file_info["path"])
    assert "CSV" in tool_out or "user_id" in tool_out

    # 4. 测试 DOCX 文件解析
    import docx
    docx_buf = io.BytesIO()
    doc = docx.Document()
    doc.add_heading("IM 群组运维与风控规范", 0)
    doc.add_paragraph("重点大群严禁发布违规涉赌涉黄链接。")
    doc.save(docx_buf)
    docx_buf.seek(0)

    docx_resp = client.post(
        "/api/agent/bots/default_assistant/upload",
        files={"file": ("guide.docx", docx_buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert docx_resp.status_code == 200
    docx_data = docx_resp.json()["file"]
    assert docx_data["type"] == "docx"
    assert "段落" in docx_data["summary"]

    docx_tool_out = read_tool.forward(docx_data["path"])
    assert "IM 群组运维与风控规范" in docx_tool_out

    # 5. 测试带文字的图片上传与智能 OCR 文本提取
    from PIL import Image, ImageDraw
    img_buf = io.BytesIO()
    test_img = Image.new("RGB", (400, 100), color=(255, 255, 255))
    d = ImageDraw.Draw(test_img)
    d.text((20, 35), "TIM Studio REST API ErrCode: 70001", fill=(0, 0, 0))
    test_img.save(img_buf, format="PNG")
    img_buf.seek(0)

    img_resp = client.post(
        "/api/agent/bots/default_assistant/upload",
        files={"file": ("error_screenshot.png", img_buf, "image/png")}
    )
    assert img_resp.status_code == 200
    img_data = img_resp.json()["file"]
    assert img_data["type"] == "png"
    assert img_data["category"] == "image"
    assert img_data["data_url"] is not None  # 验证缩略图
    assert "70001" in img_data["summary"] or "图片文件" in img_data["summary"]

    # 清理测试生成的文件
    if os.path.exists(file_info["path"]):
        os.remove(file_info["path"])
    if os.path.exists(docx_data["path"]):
        os.remove(docx_data["path"])
    if os.path.exists(img_data["path"]):
        os.remove(img_data["path"])


def test_chat_with_attachments_persistence(temp_bot_manager):
    """测试带附件提问时，附件元数据的透传与会话历史持久化"""
    from unittest.mock import patch
    from core.agent.run_manager import agent_run_manager
    client = TestClient(app)

    attachment = {
        "id": "f_12345",
        "name": "group_members.xlsx",
        "size": 2048,
        "type": "xlsx",
        "category": "table",
        "path": "/tmp/group_members.xlsx",
        "summary": "包含 50 个群成员数据",
        "preview": "列名: UserID, Role"
    }

    # 模拟以附件为载体发起会话
    with patch.object(agent_run_manager, "start_run") as mock_start:
        from core.agent.run_manager import AgentRun
        mock_run = AgentRun("run_test_att", "default_assistant", "请帮我分析群成员名单", attachments=[attachment])
        mock_run.status = "done"
        mock_start.return_value = mock_run

        resp = client.post(
            "/api/agent/bots/default_assistant/chat",
            json={
                "message": "请帮我分析群成员名单",
                "attachments": [attachment]
            }
        )
        assert resp.status_code == 200
        assert mock_start.called
        kwargs = mock_start.call_args.kwargs
        assert kwargs["attachments"] is not None
        assert kwargs["attachments"][0]["name"] == "group_members.xlsx"


def test_run_manager_memory_leak_prevention():
    """测试 AgentRunManager 具备基于容量上限和 TTL 的自动淘汰机制，防止内存泄漏"""
    from core.agent.run_manager import AgentRun, AgentRunManager

    # 创建一个最多缓存 3 个已完成 Run 的管理器
    mgr = AgentRunManager(max_cached_runs=3, run_ttl_seconds=10)

    # 1. 模拟启动 5 个任务并标记为已完成
    for i in range(5):
        run = AgentRun(f"run_{i}", "bot_1", f"msg_{i}")
        run.status = "done"
        run.end_time = time.time() - (100 if i == 0 else 0)  # run_0 模拟已过期
        mgr.runs[run.run_id] = run

    assert len(mgr.runs) == 5

    # 2. 触发淘汰机制
    with mgr._lock:
        mgr._evict_stale_runs_unlocked()

    # 3. 验证最多保留 3 个任务，且过期的 run_0 被率先淘汰
    assert len(mgr.runs) <= 3
    assert "run_0" not in mgr.runs


def test_bot_manager_orphan_file_cleanup_and_clear_history(temp_bot_manager):
    """测试会话清理时磁盘文件的同步清理以及未引用孤儿文件的主动 GC"""
    bot_id = "default_assistant"
    b_dir = temp_bot_manager._bot_dir(bot_id)
    uploads_dir = os.path.join(b_dir, "uploads")
    os.makedirs(uploads_dir, exist_ok=True)

    # 创建一个孤儿文件和一个正常引用的文件
    orphan_path = os.path.join(uploads_dir, "orphan.txt")
    referenced_path = os.path.join(uploads_dir, "referenced.txt")
    with open(orphan_path, "w") as f:
        f.write("orphan")
    with open(referenced_path, "w") as f:
        f.write("referenced")

    # 记录引用到 history
    temp_bot_manager.append_history(
        bot_id,
        "user",
        "这里有附件",
        metadata={"attachments": [{"saved_name": "referenced.txt", "path": referenced_path}]}
    )

    # 执行孤儿文件清理
    cleaned = temp_bot_manager.cleanup_unreferenced_files(bot_id)
    assert cleaned == 1
    assert not os.path.exists(orphan_path)
    assert os.path.exists(referenced_path)

    # 执行清空历史 (delete_files=True)
    temp_bot_manager.clear_history(bot_id, delete_files=True)
    assert not os.path.exists(referenced_path)


def test_tim_call_tool_high_risk_notice():
    """测试调用高危破坏性接口（解散群、删除账号）时注入安全告警标记"""
    from core.agent.tools import TimCallTool
    from core.tim_client import TimClient

    tim = TimClient(config={"sdkappid": 1400000000, "admin_user": "adm", "secret_key": "key"})
    tool = TimCallTool(tim)

    with patch.object(tim, "invoke") as mock_invoke:
        mock_invoke.return_value = {"ActionStatus": "OK", "ErrorCode": 0}

        # 1. 普通接口，无安全警告
        res_normal = tool.forward("group_open_http_svc/get_group_info", {"GroupIdList": ["@TGS#1"]})
        assert "_security_notice" not in res_normal

        # 2. 高危解散群接口，触发安全警告
        res_danger = tool.forward("group_open_http_svc/destroy_group", {"GroupId": "@TGS#1"})
        assert "_security_notice" in res_danger
        assert res_danger["_security_notice"]["is_high_risk"] is True
        assert "破坏性高危操作" in res_danger["_security_notice"]["warning"]


def test_file_hub_and_regenerate_api(temp_bot_manager):
    """测试产出物文件中心（查询/删除/压缩导出）及重新生成与截断接口"""
    client = TestClient(app)
    bot_id = "default_assistant"

    # 1. 模拟写入两个文件
    b_dir = temp_bot_manager._bot_dir(bot_id)
    uploads_dir = os.path.join(b_dir, "uploads")
    os.makedirs(uploads_dir, exist_ok=True)
    f1 = os.path.join(uploads_dir, "f_1001_report.xlsx")
    with open(f1, "w") as f:
        f.write("dummy-excel-data")

    with patch("apis.agent.bot_manager", temp_bot_manager):
        # 查询文件列表
        res_files = client.get(f"/api/agent/bots/{bot_id}/files")
        assert res_files.status_code == 200
        files_data = res_files.json()["files"]
        assert len(files_data) >= 1
        assert any(f["name"] == "f_1001_report.xlsx" for f in files_data)

        # 导出 zip
        res_export = client.get(f"/api/agent/bots/{bot_id}/files-export")
        assert res_export.status_code == 200
        assert res_export.headers["content-type"] == "application/zip"
        assert len(res_export.content) > 0

        # 模拟会话历史
        temp_bot_manager.append_history(bot_id, "user", "请输出报表")
        temp_bot_manager.append_history(bot_id, "assistant", "这是老报表内容")
        hist = temp_bot_manager.get_history(bot_id)
        assert len(hist) == 2
        user_msg_id = hist[0]["id"]

        # 测试截断历史
        res_trunc = client.post(f"/api/agent/bots/{bot_id}/truncate-history", json={"message_id": user_msg_id})
        assert res_trunc.status_code == 200
        hist_after = temp_bot_manager.get_history(bot_id)
        assert len(hist_after) == 1
        assert hist_after[0]["id"] == user_msg_id

        # 测试重新生成接口 (Mock start_run)
        from core.agent.run_manager import agent_run_manager, AgentRun
        with patch.object(agent_run_manager, "start_run") as mock_start:
            mock_run = AgentRun("run_regen_1", bot_id, "请输出报表")
            mock_run.status = "done"
            mock_start.return_value = mock_run

            res_regen = client.post(f"/api/agent/bots/{bot_id}/regenerate")
            assert res_regen.status_code == 200
            assert mock_start.called
            assert mock_start.call_args.kwargs["user_message"] == "请输出报表"

        # 删除文件
        res_del = client.delete(f"/api/agent/bots/{bot_id}/files/f_1001_report.xlsx")
        assert res_del.status_code == 200
        assert not os.path.exists(f1)


def test_bot_file_serving_and_preview(temp_bot_manager):
    """测试文件获取双路由（/bots/{id}/files/与/files/{id}/）、前缀模糊回退与inline图片预览"""
    client = TestClient(app)
    bot_id = "default_assistant"

    b_dir = temp_bot_manager._bot_dir(bot_id)
    uploads_dir = os.path.join(b_dir, "uploads")
    os.makedirs(uploads_dir, exist_ok=True)
    img_name = "f_9999999_screenshot.png"
    img_path = os.path.join(uploads_dir, img_name)
    with open(img_path, "wb") as f:
        # 写入最小有效 PNG 字节头
        f.write(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")

    with patch("apis.agent.bot_manager", temp_bot_manager):
        # 1. 测试旧版路由直接获取全名
        r1 = client.get(f"/api/agent/files/{bot_id}/{img_name}")
        assert r1.status_code == 200
        assert r1.headers["content-type"].startswith("image/png")
        assert "inline" in r1.headers.get("content-disposition", "")

        # 2. 测试新版路由带 /bots/ 前缀直接获取全名 (解决 405 问题)
        r2 = client.get(f"/api/agent/bots/{bot_id}/files/{img_name}")
        assert r2.status_code == 200
        assert r2.headers["content-type"].startswith("image/png")
        assert "inline" in r2.headers.get("content-disposition", "")

        # 3. 测试通过原文件名回退匹配前缀文件 (screenshot.png 匹配到 f_9999999_screenshot.png)
        r3 = client.get(f"/api/agent/bots/{bot_id}/files/screenshot.png")
        assert r3.status_code == 200
        assert r3.headers["content-type"].startswith("image/png")
        assert "inline" in r3.headers.get("content-disposition", "")

        # 4. 测试强制下载参数 download=true
        r4 = client.get(f"/api/agent/bots/{bot_id}/files/screenshot.png?download=true")
        assert r4.status_code == 200
        assert "attachment" in r4.headers.get("content-disposition", "")

        # 5. 测试不存在的文件返回 404
        r5 = client.get(f"/api/agent/bots/{bot_id}/files/not_found.png")
        assert r5.status_code == 404

        # 6. 测试删除支持原名回退
        r6 = client.delete(f"/api/agent/bots/{bot_id}/files/screenshot.png")
        assert r6.status_code == 200
        assert not os.path.exists(img_path)


def test_render_chart_helper():
    """测试 ECharts 辅助输出函数"""
    from core.agent.engine import render_chart
    chart_md = render_chart(
        chart_type="bar",
        title="近7天违规消息分布",
        labels=["周一", "周二", "周三"],
        data=[12, 19, 5]
    )
    assert "```echarts" in chart_md
    assert "近7天违规消息分布" in chart_md
    assert '"type": "bar"' in chart_md
    assert '"data": [\n    12,\n    19,\n    5\n  ]' in chart_md or '"data": [' in chart_md


def test_rag_manager_and_knowledge_api(temp_bot_manager, tmp_path):
    """测试轻量级知识库与文档切片、BM25 检索及 RAG 接口"""
    from core.agent.rag_manager import RagManager
    rag = RagManager(data_root=str(tmp_path))
    bot_id = "test_rag_bot"
    temp_bot_manager.create_bot(bot_id=bot_id, name="测试知识库助手", avatar="📚", description="测试知识库")

    # 1. 上传 Markdown 规范文档
    md_content = """# 腾讯云 IM 错误码处置规范

## 错误码 10007 排查 SOP
当遇到 10007 错误码时，表示目标群组不存在或已被解散。
排查步骤：
1. 确认 GroupId 是否带有 @TGS# 前缀；
2. 调用 get_group_info 验证群组当前生命周期状态；
3. 若群组已解散，记录审计日志并提示业务层。

## 解散群组安全操作准则
解散群属于破坏性高危动作，必须向运维或业务负责人双人核准后方可执行。
严禁未经二次确认大批量清退生产群。
""".encode("utf-8")

    doc_meta = rag.add_document(bot_id, "IM_error_sop.md", md_content, tags=["运维规范"])
    assert doc_meta["name"] == "IM_error_sop.md"
    assert doc_meta["chunks_count"] >= 2
    assert doc_meta["enabled"] is True

    # 2. 测试 BM25 检索
    hits = rag.search(bot_id, "错误码 10007 怎么排查")
    assert len(hits) > 0
    assert "10007" in hits[0]["content"] or "10007" in hits[0]["title"]
    assert hits[0]["score"] > 0

    hits_danger = rag.search(bot_id, "解散群组有什么要求")
    assert len(hits_danger) > 0
    assert "高危" in hits_danger[0]["content"]

    # 3. 测试启用/禁用切换
    rag.toggle_document(bot_id, doc_meta["id"], enabled=False)
    hits_disabled = rag.search(bot_id, "错误码 10007")
    assert len(hits_disabled) == 0

    rag.toggle_document(bot_id, doc_meta["id"], enabled=True)
    hits_enabled = rag.search(bot_id, "错误码 10007")
    assert len(hits_enabled) > 0

    # 4. 测试 API 接口
    client = TestClient(app)
    with patch("apis.agent.bot_manager", temp_bot_manager), patch("apis.agent.rag_manager", rag):
        # 4.1 列出文档
        res_list = client.get(f"/api/agent/bots/{bot_id}/knowledge/docs")
        assert res_list.status_code == 200
        docs = res_list.json()["docs"]
        assert len(docs) == 1
        assert docs[0]["name"] == "IM_error_sop.md"

        # 4.2 在线搜索测试
        res_search = client.post(
            f"/api/agent/bots/{bot_id}/knowledge/search",
            json={"query": "10007", "top_k": 3}
        )
        assert res_search.status_code == 200
        assert res_search.json()["count"] >= 1

        # 4.3 切换禁用
        res_toggle = client.put(
            f"/api/agent/bots/{bot_id}/knowledge/docs/{doc_meta['id']}",
            json={"enabled": False}
        )
        assert res_toggle.status_code == 200
        assert rag.get_document(bot_id, doc_meta["id"])["enabled"] is False

        # 4.4 删除文档
        res_del = client.delete(f"/api/agent/bots/{bot_id}/knowledge/docs/{doc_meta['id']}")
        assert res_del.status_code == 200
        assert len(rag.list_documents(bot_id)) == 0


def test_orchestration_topology_api(temp_bot_manager):
    """测试多智能体协同流水记录与全局网络拓扑 API"""
    client = TestClient(app)

    # 记录两次协同事件
    event1 = {
        "caller_bot_id": "default_assistant",
        "caller_bot_name": "TIM 全能助手",
        "target_bot_id": "risk_inspector",
        "target_bot_name": "社群风控官",
        "target_avatar": "🛡️",
        "task": "排查敏感消息并批量禁言违规成员",
        "status": "success",
        "result_preview": "已成功处理 3 个违规成员禁言",
        "elapsed_seconds": 1.25,
        "timestamp": "14:30:00"
    }
    temp_bot_manager.record_orchestration(event1)

    with patch("apis.agent.bot_manager", temp_bot_manager):
        # 获取全局拓扑
        res = client.get("/api/agent/topology")
        assert res.status_code == 200
        data = res.json()
        assert "nodes" in data
        assert "links" in data
        assert any(l["source"] == "default_assistant" and l["target"] == "risk_inspector" for l in data["links"])

        # 获取单 BOT 协同流水
        res_bot = client.get("/api/agent/bots/default_assistant/orchestrations")
        assert res_bot.status_code == 200
        bot_events = res_bot.json()["orchestrations"]
        assert len(bot_events) >= 1
        assert bot_events[0]["target_bot_name"] == "社群风控官"


def test_agent_robustness_and_concurrency(temp_bot_manager):
    """测试强化修复：技能执行超时保护、RAG 并发安全锁、Regenerate 无重复历史消息"""
    from unittest.mock import MagicMock, patch
    from core.agent.tools import RunSkillTool
    from core.agent.rag_manager import rag_manager
    from core.tim_client import TimClient
    import concurrent.futures

    bot_id = "default_assistant"

    # 1. 验证 RunSkillTool 对死循环/超时的防护
    temp_bot_manager.save_skill(
        bot_id=bot_id,
        name="infinite_loop_skill",
        description="测试超时",
        code="import time\ntime.sleep(30)\n"
    )
    runner = RunSkillTool(bot_id, temp_bot_manager, TimClient())
    with patch("concurrent.futures.ThreadPoolExecutor.submit") as mock_submit:
        # 模拟超时
        mock_future = MagicMock()
        mock_future.result.side_effect = concurrent.futures.TimeoutError()
        mock_submit.return_value = mock_future
        res = runner.forward("infinite_loop_skill")
        assert res["success"] is False
        assert "超时" in res["error"]

    # 2. 验证 RAG 多线程并发文件锁
    def _concurrent_add(i):
        return rag_manager.add_document(bot_id, f"doc_{i}.txt", f"这是测试文档内容 {i}".encode("utf-8"))

    with concurrent.futures.ThreadPoolExecutor(max_workers=5) as executor:
        futures = [executor.submit(_concurrent_add, i) for i in range(5)]
        results = [f.result() for f in futures]
        assert len(results) == 5

    docs = rag_manager.list_documents(bot_id)
    assert len(docs) >= 5


def test_consent_gate_and_chat_export_api(temp_bot_manager):
    """测试高危人机协同门禁 (Consent Gate) 与会话排障记录导出 (Markdown/HTML)"""
    client = TestClient(app)
    bot_id = "default_assistant"

    # 1. 模拟历史会话并测试导出
    temp_bot_manager.append_history(
        bot_id=bot_id,
        role="user",
        content="请帮我排查群组活跃度与违规敏感消息",
        metadata={"attachments": [{"name": "test.csv", "type": "csv", "size_formatted": "12 KB"}]}
    )
    temp_bot_manager.append_history(
        bot_id=bot_id,
        role="assistant",
        content="已完成排查，未发现严重违规发言。",
        metadata={
            "thought": "首先调用群组消息拉取接口分析...",
            "code_blocks": [{"step": 1, "code": "print('testing')"}],
            "execution_logs": [{"step": 1, "output": "testing"}]
        }
    )

    with patch("apis.agent.bot_manager", temp_bot_manager):
        # 导出 Markdown
        res_md = client.get(f"/api/agent/bots/{bot_id}/export?format=markdown")
        assert res_md.status_code == 200
        assert "text/markdown" in res_md.headers["content-type"]
        assert "排障分析报告" in res_md.text
        assert "test.csv" in res_md.text

        # 导出 HTML
        res_html = client.get(f"/api/agent/bots/{bot_id}/export?format=html")
        assert res_html.status_code == 200
        assert "text/html" in res_html.headers["content-type"]
        assert "<!DOCTYPE html>" in res_html.text
        assert "TIM 智能体排障分析报告" in res_html.text

        # 2. 测试高危门禁 TimCallTool 触发审批与用户决策
        from core.agent.tools import TimCallTool
        from core.tim_client import TimClient
        from unittest.mock import MagicMock

        # 模拟待审批请求回调
        approval_captured = {}
        def mock_approval_requester(appr_data):
            approval_captured.update(appr_data)
            return "reject"  # 用户拒绝

        tim_tool = TimCallTool(TimClient(), approval_requester=mock_approval_requester)
        # 触发高危接口 destroy_group
        res_destroy = tim_tool.forward("group_open_http_svc/destroy_group", {"GroupId": "@TGS#test_group"})
        assert res_destroy["ActionStatus"] == "CANCELLED"
        assert res_destroy["ErrorCode"] == -999
        assert "拒绝" in res_destroy["ErrorInfo"]
        assert approval_captured.get("api_path") == "group_open_http_svc/destroy_group"
        assert "@TGS#test_group" in approval_captured.get("target_summary", "")


def test_provider_ssot_templates_and_security_endpoints(temp_bot_manager):
    """测试大模型提供商单点真实源 (SSoT)、预设角色模板库以及高危操作安全拦截配置"""
    client = TestClient(app)

    with patch("apis.agent.bot_manager", temp_bot_manager):
        # 1. 验证大模型配置提供商包含主流国内厂商与标准预设
        res_cfg = client.get("/api/agent/config")
        assert res_cfg.status_code == 200
        providers = res_cfg.json()["providers"]
        provider_ids = [p["id"] for p in providers]
        assert "deepseek" in provider_ids
        assert "qwen" in provider_ids
        assert "moonshot" in provider_ids
        assert "siliconflow" in provider_ids
        assert "zhipu" in provider_ids
        assert "custom" in provider_ids

        # 2. 验证预设模板库返回
        res_tmpl = client.get("/api/agent/templates")
        assert res_tmpl.status_code == 200
        templates = res_tmpl.json()["templates"]
        assert len(templates) >= 3
        tmpl_names = [t["name"] for t in templates]
        assert "社群风控官" in tmpl_names
        assert "消息巡检员" in tmpl_names
        assert "运维保障官" in tmpl_names

        # 3. 验证高危接口拦截规则 GET 与 POST
        res_sec = client.get("/api/agent/security/high-risk-keywords")
        assert res_sec.status_code == 200
        data_sec = res_sec.json()
        assert "destroy_group" in data_sec["keywords"]

        # 更新高危关键词
        new_kws = ["destroy_group", "account_delete", "custom_drop_table"]
        res_post_sec = client.post("/api/agent/security/high-risk-keywords", json={"keywords": new_kws})
        assert res_post_sec.status_code == 200
        assert "custom_drop_table" in res_post_sec.json()["keywords"]
        assert temp_bot_manager.get_high_risk_keywords() == ["destroy_group", "account_delete", "custom_drop_table"]


def test_fuzzy_memory_recall_and_sandbox_file_tools(temp_bot_manager):
    """测试长期记忆自然语言模糊召回与沙箱资产文件工具"""
    from core.agent.tools import RecallMemoryTool

    # 1. 写入测试记忆
    bot_id = "mem_test_bot"
    temp_bot_manager.create_bot(bot_id=bot_id, name="记忆测试者")
    temp_bot_manager.save_memory(
        bot_id=bot_id,
        key="group_security_whitelist",
        content="白名单群组名单包括 @TGS#alpha 和 @TGS#beta，这些群严禁执行自动解散"
    )

    recall_tool = RecallMemoryTool(bot_id, temp_bot_manager)
    # 自然语言模糊查询（包含多个词素）
    res = recall_tool.forward("查询当前系统的白名单群组与安全规则")
    assert len(res) == 1
    assert res[0]["key"] == "group_security_whitelist"
    assert "@TGS#alpha" in res[0]["content"]

    # 2. 验证沙箱产出物目录与文件写入 helper
    from core.agent.engine import agent_engine
    bot_uploads = os.path.join(temp_bot_manager._bot_dir(bot_id), "uploads")
    os.makedirs(bot_uploads, exist_ok=True)

    test_file_path = os.path.join(bot_uploads, "sandbox_output.txt")
    with open(test_file_path, "w", encoding="utf-8") as f:
        f.write("sandbox test content")

    files = temp_bot_manager.list_bot_files(bot_id)
    file_names = [f["name"] for f in files]
    assert "sandbox_output.txt" in file_names


def test_fetch_available_models_api():
    """测试通过 /api/agent/models 接口在线拉取与降级预设模型标识"""
    from unittest.mock import patch, MagicMock

    client = TestClient(app)

    # 1. 模拟 Base URL 为空时，返回特定提供商的推荐模型预设
    res = client.post("/api/agent/models", json={"provider": "deepseek", "base_url": ""})
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "deepseek-chat" in data["models"]
    assert data["source"] == "preset"

    # 2. 模拟远程 OpenAI 兼容 /models 成功响应
    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.json.return_value = {
        "object": "list",
        "data": [
            {"id": "qwen2.5-coder:32b", "object": "model"},
            {"id": "custom-tuned-v1", "object": "model"}
        ]
    }
    with patch("requests.get", return_value=mock_resp):
        res2 = client.post("/api/agent/models", json={
            "provider": "custom",
            "base_url": "http://mock-llm.local/v1",
            "api_key": "sk-mock-key"
        })
        assert res2.status_code == 200
        data2 = res2.json()
        assert data2["success"] is True
        assert data2["source"] == "remote"
        assert "qwen2.5-coder:32b" in data2["models"]
        assert "custom-tuned-v1" in data2["models"]

    # 3. 模拟远程超时/异常时的优雅降级
    with patch("requests.get", side_effect=Exception("Connection timed out")):
        res3 = client.post("/api/agent/models", json={
            "provider": "openai",
            "base_url": "https://api.openai.com/v1",
            "api_key": "sk-invalid"
        })
        assert res3.status_code == 200
        data3 = res3.json()
        assert data3["success"] is False
        assert data3["source"] == "fallback"
        assert "gpt-4o" in data3["models"]


def test_clean_final_answer_sanitization():
    """测试 final_answer 的净化能力：剥离 Calling tools 内部协议、DSML 标签，规范化 code 标签"""
    from core.agent.engine import clean_final_answer

    # 1. 包含 Calling tools 泄露
    dirty_text1 = (
        "已为您获取前两个群。\n"
        "Calling tools:\n"
        "[{'id': 'call_11', 'type': 'function', 'function': {'name': 'python_interpreter', 'arguments': 'print(1)'}}]\n"
        "后续分析如下。"
    )
    cleaned1, thought1 = clean_final_answer(dirty_text1)
    assert "Calling tools" not in cleaned1
    assert "call_11" not in cleaned1
    assert "已为您获取前两个群" in cleaned1
    assert "后续分析如下" in cleaned1

    # 2. 包含 <code>...</code> 标签
    dirty_text2 = (
        "分析如下：\n"
        "<code>\n"
        "import json\n"
        "print('hello')\n"
        "</code>"
    )
    cleaned2, thought2 = clean_final_answer(dirty_text2)
    assert "<code>" not in cleaned2
    assert "</code>" not in cleaned2
    assert "```python" in cleaned2
    assert "import json" in cleaned2

    # 3. 包含 DSML 原生调用标签
    dirty_text3 = (
        "数据处理完成。\n"
        "<｜｜DSML｜｜ calls>\n"
        "<｜｜DSML｜｜ invoke name=\"python_interpreter\">\n"
        "<parameter name=\"code\">print('secret')</parameter>\n"
        "</invoke>"
    )
    cleaned3, thought3 = clean_final_answer(dirty_text3)
    assert "<｜｜DSML｜｜" not in cleaned3
    assert "数据处理完成" in cleaned3


def test_robust_parse_code_blobs():
    """测试增强型代码块解析器支持 standard code tags、DSML 与 XML parameter"""
    from core.agent.engine import robust_parse_code_blobs

    # 1. 常见 <code> 标签
    code1 = robust_parse_code_blobs("Thought: test\n<code>\nprint('code1')\n</code>", ("<code>", "</code>"))
    assert "print('code1')" in code1

    # 2. Markdown 块
    code2 = robust_parse_code_blobs("Thought: test\n```python\nprint('code2')\n```", ("<code>", "</code>"))
    assert "print('code2')" in code2

    # 3. DeepSeek DSML parameter code
    dsml_text = (
        "Let me do it.<｜｜DSML｜｜ calls>\n"
        "<｜｜DSML｜｜ invoke name=\"python_interpreter\">\n"
        "<parameter name=\"code\">print('code_dsml')</parameter>\n"
        "</invoke>"
    )
    code3 = robust_parse_code_blobs(dsml_text, ("<code>", "</code>"))
    assert "print('code_dsml')" in code3

    # 4. XML python 标签
    xml_text = "Thought: xml\n<python>\nprint('code_xml')\n</python>"
    code4 = robust_parse_code_blobs(xml_text, ("<code>", "</code>"))
    assert "print('code_xml')" in code4


def test_sandbox_state_persistence_and_clearing(temp_bot_manager):
    """测试跨轮对话沙箱状态持久化与清空"""
    from core.agent.engine import agent_engine

    bot_id = "test_sandbox_bot"
    agent_engine._bot_sandbox_states[bot_id] = {
        "top5": [{"gid": "123", "name": "群A"}],
        "total_count": 42
    }

    assert bot_id in agent_engine._bot_sandbox_states
    assert agent_engine._bot_sandbox_states[bot_id]["total_count"] == 42

    # 清空沙箱
    agent_engine.clear_sandbox(bot_id)
    assert bot_id not in agent_engine._bot_sandbox_states


def test_hierarchical_memory_manager_all_four_layers(temp_bot_manager, tmp_path):
    """综合测试四层分级记忆与无限上下文架构 (Layer 1~4)"""
    import json
    import os
    from core.agent.memory import HierarchicalMemoryManager
    from core.agent.models import BotConfig
    import pandas as pd

    mem_mgr = HierarchicalMemoryManager()
    bot_id = "test_memory_bot"
    bot_dir = str(tmp_path / bot_id)
    os.makedirs(bot_dir, exist_ok=True)

    # 1. Layer 4: 沙箱运行时变量自省 (DataFrame, list, dict, primitives)
    df_sample = pd.DataFrame([{"GroupId": "GROUP_CHAT_V2:974", "Members": 100}])
    sandbox_state = {
        "df_result": df_sample,
        "target_groups": ["GROUP_CHAT_V2:974", "GROUP_CHAT_V2:1043"],
        "summary_dict": {"total": 2, "active": 1},
        "query_count": 5
    }
    sandbox_prompt, var_map = mem_mgr.inspect_sandbox_variables(sandbox_state)
    assert "df_result" in sandbox_prompt
    assert "pandas.DataFrame" in sandbox_prompt
    assert "target_groups" in sandbox_prompt
    assert "query_count" in var_map
    assert var_map["query_count"] == "int = 5"

    # 2. Layer 3: 实体提取 (从文本及代码块中抓取群号和接口)
    messages = [
        {"role": "user", "content": "请拉取 @TGS#abc12345 和 GROUP_CHAT_V2:974 的消息"},
        {
            "role": "assistant",
            "content": "正在调用 tim_call('group_open_http_svc/group_msg_get_simple')",
            "metadata": {
                "code_blocks": [{"code": "tim_call('group_open_http_svc/group_msg_get_simple', {'GroupId': 'LIVE_999'})"}]
            }
        }
    ]
    entities = mem_mgr.extract_entities_from_messages(messages)
    assert "@TGS#abc12345" in entities["groups"]
    assert "GROUP_CHAT_V2:974" in entities["groups"]
    assert "LIVE_999" in entities["groups"]
    assert "group_open_http_svc/group_msg_get_simple" in entities["apis"]

    # 3. Layer 1: 即时工作记忆与观测蒸馏 (Observation Distillation)
    distilled = mem_mgr.distill_assistant_turn(
        "成功拉取消息数据并完成清洗。",
        {"code_blocks": [{"code": "print(1)"}], "tool_calls": [{"name": "tim_call"}]},
        max_len=100
    )
    assert "成功拉取消息数据" in distilled
    assert "执行了 1 段 Python 脚本" in distilled
    assert "tim_call" in distilled

    working_memory_turns = mem_mgr.build_working_memory(messages)
    assert len(working_memory_turns) == 2
    assert "用户:" in working_memory_turns[0]
    assert "助手:" in working_memory_turns[1]

    # 4. Layer 2: 会话演化滚动摘要增量折叠 (Running Session Summary)
    older_history = []
    for i in range(25):
        older_history.append({"id": f"msg-{i}", "role": "user", "content": f"第 {i} 轮分析 GROUP_CHAT_V2:{1000 + i}"})
        older_history.append({"id": f"ast-{i}", "role": "assistant", "content": f"已完成第 {i} 轮 group_open_http_svc/group_msg_get_simple 数据统计"})

    summary_text = mem_mgr.build_or_update_session_summary(bot_id, bot_dir, older_history)
    assert "【会话至今已累计推进" in summary_text
    assert "核心业务脉络与阶段目标" in summary_text
    assert "阶段性执行成果与关键发现" in summary_text

    # 验证 session_summary.json 是否已持久化
    summary_file = os.path.join(bot_dir, "session_summary.json")
    assert os.path.exists(summary_file)
    with open(summary_file, "r", encoding="utf-8") as f:
        data = json.load(f)
        assert data["total_compacted_count"] == 50
        assert data["last_compacted_msg_id"] == "ast-24"

    # 5. 清理记忆
    mem_mgr.clear_session_memory(bot_id, bot_dir)
    assert not os.path.exists(summary_file)


def test_context_budget_dynamic_configuration(temp_bot_manager, tmp_path):
    """测试即时工作记忆动态预算配置 (Context Budgeting)"""
    from core.agent.memory import HierarchicalMemoryManager
    from core.agent.models import BotConfig, LLMConfig

    mem_mgr = HierarchicalMemoryManager()
    bot_id = "test_budget_bot"
    bot_dir = str(tmp_path / bot_id)
    os.makedirs(bot_dir, exist_ok=True)

    bot = BotConfig(
        id=bot_id,
        name="测试预算Bot",
        context_budget=12000,
        max_context_turns=10
    )

    # 构造 15 轮长消息历史
    history = []
    for i in range(15):
        history.append({"id": f"u-{i}", "role": "user", "content": f"用户提问 {i}: " + "长文本内容" * 50})
        history.append({"id": f"a-{i}", "role": "assistant", "content": f"助手回答 {i}: " + "详细排查结果" * 50})

    # 1. 采用 bot 自定义小预算：仅保留 10 轮且字符不超过 12000
    ctx_small = mem_mgr.build_hierarchical_context(
        bot_id=bot_id,
        bot=bot,
        user_message="当前最新需求",
        all_history=history,
        saved_state={},
        attachments=[],
        system_prompt="System Prompt",
        manager=temp_bot_manager,
        uploads_dir=bot_dir,
        context_budget=12000,
        max_context_turns=10
    )
    assert ctx_small.working_memory_turns <= 10
    assert ctx_small.compacted_history_count == 20  # 30 条消息中前 20 条被判定为超期

    # 2. 采用超大模型预算：64000 字符，保留 25 轮
    ctx_large = mem_mgr.build_hierarchical_context(
        bot_id=bot_id,
        bot=bot,
        user_message="当前最新需求",
        all_history=history,
        saved_state={},
        attachments=[],
        system_prompt="System Prompt",
        manager=temp_bot_manager,
        uploads_dir=bot_dir,
        context_budget=64000,
        max_context_turns=25
    )
    assert ctx_large.working_memory_turns > ctx_small.working_memory_turns














